"""Score the two completed annotation workbooks against each other and the heuristic.

Reports Cohen's kappa between annotators (Q1 damage, Q3 note) and, on samples
where both annotators agree, the heuristic's accuracy for "did damage occur"
(Q1), "where does damage start" (Q2), and "is the detected note real" (Q3).

  .venv/bin/python kappa.py --a annotation/annotation_annotator_A.xlsx \
      --b annotation/annotation_annotator_B.xlsx --key annotation/KEY_do_not_share.csv \
      --out annotation/agreement.txt
"""
import argparse
import re

import pandas as pd
from openpyxl import load_workbook


def answers(path):
    ws = load_workbook(path, data_only=True)['Annotate']
    rows = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if isinstance(r[0], str) and re.fullmatch(r'S\d{3}', r[0]):
            rows[r[0]] = {'q1': r[6], 'q2': r[7], 'q3': r[8]}
    return pd.DataFrame.from_dict(rows, orient='index')


def kappa(x, y):
    m = x.notna() & y.notna()
    x, y = x[m].astype(str), y[m].astype(str)
    if len(x) == 0:
        return float('nan'), 0
    po = (x == y).mean()
    cats = set(x) | set(y)
    pe = sum((x == c).mean() * (y == c).mean() for c in cats)
    return (po - pe) / (1 - pe) if pe < 1 else 1.0, len(x)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--a', required=True)
    ap.add_argument('--b', required=True)
    ap.add_argument('--key', required=True)
    ap.add_argument('--out', required=True)
    a = ap.parse_args()
    A, B = answers(a.a), answers(a.b)
    key = pd.read_csv(a.key, dtype=str).set_index('id')
    lines = []
    p = lambda s='': (print(s), lines.append(s))

    for q, name in (('q1', 'Q1 damage'), ('q3', 'Q3 note real')):
        x, y = A[q], B[q]
        if q == 'q3':
            keep = (x != 'N/A') & (y != 'N/A')
            x, y = x[keep], y[keep]
        k, n = kappa(x, y)
        p(f'{name:14s} Cohen kappa = {k:.3f}  (n = {n}, raw agreement {(x.astype(str) == y.astype(str)).mean():.1%})')

    both = A.join(B, lsuffix='_a', rsuffix='_b').join(key)
    agree1 = both[(both.q1_a == both.q1_b) & both.q1_a.isin(['Yes', 'No'])]
    acc = (agree1.q1_a == agree1.heuristic_damaged).mean()
    p(f'\nHeuristic vs. agreed Q1 (n = {len(agree1)}): accuracy {acc:.1%}')
    p(pd.crosstab(agree1.heuristic_damaged, agree1.q1_a, rownames=['heuristic'], colnames=['annotators']).to_string())
    a2 = agree1[(agree1.q1_a == 'Yes') & (agree1.heuristic_damaged == 'Yes')
                & (agree1.q2_a.astype(str) == agree1.q2_b.astype(str))]
    p(f'\nQ2 first damage event, annotators agree (n = {len(a2)}): heuristic identical in '
      f'{(a2.q2_a.astype(str) == a2.heuristic_first_event.astype(str)).mean():.1%}')
    a3 = both[(both.q3_a == both.q3_b) & both.q3_a.isin(['Yes', 'No']) & (both.note_detected == 'Yes')]
    p(f'Q3 detected notes judged real (n = {len(a3)}): {(a3.q3_a == "Yes").mean():.1%}')
    open(a.out, 'w').write('\n'.join(lines) + '\n')


if __name__ == '__main__':
    main()
