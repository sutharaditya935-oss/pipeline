"""Sort the benign candidate list by whether it can run under the lab's default UAC.

The lab launches samples unelevated (Explorer's filtered token), so any program
that demands administrator rights fails with ERROR_ELEVATION_REQUIRED (740), as 68
ransomware analyses did. Groups (an ESTIMATE from how each tool is normally
packaged; confirm with the launch test in benign/README.md before relying on it):

  A  likely runs unelevated: portable zip / standalone or static exe / built-in tool
  B  uncertain: has an installer that may or may not install per-user, or needs a
     portable build we must source, or depends on other software; verify
  C  likely needs administrator (installs to Program Files / all users, drivers,
     services, volume-level access) or cannot work offline: exclude from the
     default-UAC baseline

  python benign/classify_corpus.py --xlsx ~/Downloads/benign_corpus.xlsx \
      --out benign/batch1/sorted_candidates
"""
import argparse, csv, os

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# row number (column "#") -> (group, reason)
G = {}
def put(grp, rows, why):
    for r in rows:
        G[r] = (grp, why)

put('A', [1, 2, 3, 4, 51], 'use the standalone/portable 7-Zip build (7za/7z), not the installer')
put('A', [5, 6], 'PeaZip portable')
put('A', [13, 14, 15, 16], 'standalone command-line binary')
put('A', [21, 22, 23], 'static single executable')
put('A', [28], 'portable build, runs per-user')
put('A', [29, 30, 122], 'built-in Windows tool, no admin needed for user folders')
put('A', [31], 'cwRsync/MSYS2 binaries copied, no install')
put('A', [44, 45, 46, 47], 'standalone binary / portable OpenSSL')
put('A', [49], 'cipher /e on a user folder works unelevated (verify EFS on this edition)')
put('A', [52], 'Picocrypt portable exe')
put('A', [57], '7-Zip SFX archive extracts to a user folder')
put('A', [71, 73, 76], 'browser portable/zip build')
put('A', [80, 81], 'portable build')
put('A', [85, 86], 'KeePass/KeePassXC portable')
put('A', [88], 'Calibre Portable')
put('A', [90], 'Inkscape portable archive')
put('A', [91, 92, 104], 'PortableGit')
put('A', [94, 95, 97, 99, 103], 'portable/zip toolchain, no installer (needs matching compiler/runtime zip)')
put('A', [105, 106, 107, 108, 110, 111, 112, 113, 114], 'portable/static build')
put('A', [118, 120, 123, 126], 'portable tool (Sysinternals sigcheck/du, WinDirStat, BleachBit, SDelete)')
put('A', [131, 136, 138, 139], 'portable/static command-line tool')

put('B', [7, 8, 9], 'WinRAR installer wants admin; need a portable/rar.exe build')
put('B', [10, 11], 'Bandizip has a portable edition; installer wants admin')
put('B', [12], 'MSIX package; per-user sideload may work, verify')
put('B', [18], 'IZArc installer; check for portable zip')
put('B', [19, 20], 'Duplicati: use the portable zip instead of the MSI')
put('B', [26, 27], 'FreeFileSync supports a portable install')
put('B', [35, 36], 'installer; check for portable mode')
put('B', [37], 'VeraCrypt portable extraction can create containers; mounting needs a driver (see C)')
put('B', [39, 40], 'use GnuPG portable binaries, not the Gpg4win installer')
put('B', [43, 53], 'per-user install possible; verify')
put('B', [50], 'cipher /w on a user volume; may need elevation')
put('B', [54, 55], 'depends on the specific installer (per-user vs Program Files)')
put('B', [60], 'Git installer defaults to Program Files; per-user option exists')
put('B', [69, 70, 72, 74, 75], 'no official portable build; per-user install possible, verify')
put('B', [77, 78, 79], 'office suite needs a portable build (e.g. a portable edition)')
put('B', [83, 84, 87], 'per-user installer or portable archive, verify')
put('B', [93], 'MSYS2 base archive can be extracted without install; verify')
put('B', [96, 98, 101, 102], 'needs other software (compiler/JDK/Python) staged first')
put('B', [109], 'OBS portable mode exists; recording needs graphics, verify offline')
put('B', [116, 117], 'MFT mode needs admin; falls back to a slower scan unelevated, verify')
put('B', [119, 121], 'installer/shell extension; find a portable alternative')
put('B', [140], 'installer; verify per-user')

put('C', [17], 'WinZip installer needs admin')
put('C', [24], 'needs WSL, not in the guest')
put('C', [25, 32, 33, 34], 'installs services/drivers and needs admin (Veeam, UrBackup, Macrium, Cobian)')
put('C', [38, 41, 42], 'volume mount driver / shell extension / installer needs admin')
put('C', [48, 128, 129], 'volume-level or servicing command that requires administrator')
put('C', [56, 58, 59, 61, 62, 63, 64, 65, 66, 67, 68], 'installs to Program Files / all users (MSI or admin installer)')
put('C', [82, 89, 115], 'installer only, needs admin')
put('C', [100], 'Docker CLI needs a running daemon, not available')
put('C', [124, 125, 127, 130, 133, 135, 141, 137], 'installer, service or driver needs admin (Eraser, CCleaner, Recuva, Procmon, Malwarebytes, Sophos, Seafile, Nextcloud)')
put('C', [132], 'Defender scan is not meaningful in this guest (C:\\ is excluded)')
put('C', [134], 'ESET online scanner needs internet')

FILL = {'A': 'E2EFDA', 'B': 'FFF2CC', 'C': 'F8CBAD'}
TITLE = {'A': 'A likely OK', 'B': 'B verify', 'C': 'C needs admin or excluded'}


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--xlsx', required=True)
    ap.add_argument('--out', required=True, help='output path without extension')
    a = ap.parse_args()
    ws = load_workbook(a.xlsx, data_only=True)['Corpus']
    head = [c.value for c in ws[1]]
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]
    missing = [int(r[0]) for r in rows if int(r[0]) not in G]
    extra = [k for k in G if k not in {int(r[0]) for r in rows}]
    if missing or extra:
        raise SystemExit(f'unclassified rows: {missing}; unknown rows: {extra}')

    out = [{'#': int(r[0]), 'group': G[int(r[0])][0], 'reason': G[int(r[0])][1],
            **{h: r[i] for i, h in enumerate(head) if h != '#'}} for r in rows]
    out.sort(key=lambda d: (d['group'], d['Category'], d['#']))
    os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
    with open(a.out + '.csv', 'w', newline='') as f:
        w = csv.DictWriter(f, fieldnames=list(out[0])); w.writeheader(); w.writerows(out)

    wb = Workbook(); wb.remove(wb.active)
    cols = ['#', 'group', 'reason'] + [h for h in head if h != '#']
    for g in 'ABC':
        sh = wb.create_sheet(TITLE[g])
        sh.append(cols)
        for d in (d for d in out if d['group'] == g):
            sh.append([d[c] for c in cols])
        for c in sh[1]:
            c.font = Font(name='Arial', bold=True, size=9); c.fill = PatternFill('solid', fgColor=FILL[g])
            c.alignment = Alignment(wrap_text=True, vertical='center')
        for row in sh.iter_rows(min_row=2):
            for c in row:
                c.font = Font(name='Arial', size=9)
        for i, w_ in enumerate([5, 7, 46, 22, 26, 22, 38, 20, 40, 10, 22, 12, 12, 16][:len(cols)], 1):
            sh.column_dimensions[get_column_letter(i)].width = w_
        sh.freeze_panes = 'D2'
    note = wb.create_sheet('Read me', 0)
    for line in ['Sorted by whether a program can run under the lab\'s default UAC (unelevated launch).',
                 'This sort is an ESTIMATE from how each tool is normally packaged. Confirm each program with the',
                 'launch test (benign/README.md) before it enters the baseline.',
                 'A = likely OK (portable/standalone/built-in)   B = verify   C = likely needs admin or cannot run offline.',
                 'Tasks from the same program share a download; fill SHA-256/version/source for each program you keep.']:
        note.append([line])
    note.column_dimensions['A'].width = 120
    wb.save(a.out + '.xlsx')
    n = {g: sum(d['group'] == g for d in out) for g in 'ABC'}
    progs = {g: len({d['Program'] for d in out if d['group'] == g}) for g in 'ABC'}
    print(f'tasks: A={n["A"]} B={n["B"]} C={n["C"]} (total {len(out)}); distinct programs: A={progs["A"]} B={progs["B"]} C={progs["C"]}')
    for cat in sorted({d['Category'] for d in out}):
        c = {g: sum(d['group'] == g and d['Category'] == cat for d in out) for g in 'ABC'}
        print(f'  {cat:28s} A={c["A"]:2d} B={c["B"]:2d} C={c["C"]:2d}')


if __name__ == '__main__':
    main()
